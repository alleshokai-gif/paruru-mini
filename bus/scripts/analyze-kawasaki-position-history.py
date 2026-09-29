"""Read-only, aggregate-only coverage from a temporary observation workbook export."""

import collections
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main(workbook_path, static_path):
    position = json.loads(Path(static_path).read_text(encoding="utf-8"))
    trips = position["trips"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = workbook["Bus_Observation_Raw"].iter_rows(values_only=True)
    header = next(rows)
    expected = ("observation_id", "service_date", "direction_id", "route_id", "trip_id",
                "vehicle_hash", "gps_timestamp", "position_lat", "position_lon")
    if not set(expected).issubset(header):
        raise ValueError("OBSERVATION_HEADER_MISMATCH")
    col = {name: header.index(name) for name in expected}
    chain_stats = collections.defaultdict(lambda: {"rows": 0, "gps": 0, "days": set(),
                                                    "tripDays": set(), "traces": set()})
    query_stats = collections.Counter()
    missing_trips = 0
    for row in rows:
        if not row or not row[col["observation_id"]]:
            continue
        direction = str(row[col["direction_id"]])
        query_stats[direction] += 1
        trip_id = str(row[col["trip_id"]])
        trip = trips.get(trip_id)
        if not trip:
            missing_trips += 1
            continue
        chain_id = trip["chainId"]
        stats = chain_stats[chain_id]
        stats["rows"] += 1
        date = str(row[col["service_date"]])
        stats["days"].add(date)
        stats["tripDays"].add((date, trip_id))
        vehicle = str(row[col["vehicle_hash"]] or "")
        if vehicle.startswith("veh_"):
            stats["traces"].add((date, trip_id, vehicle))
        if row[col["position_lat"]] is not None and row[col["position_lon"]] is not None:
            stats["gps"] += 1
    report = {"status": "KAWASAKI_POSITION_HISTORY", "rawRows": sum(query_stats.values()),
              "rawDirections": dict(query_stats), "missingSidecarRows": missing_trips,
              "chains": [{"chainId": chain_id, "rows": stats["rows"], "gpsRows": stats["gps"],
                          "serviceDays": len(stats["days"]), "tripDays": len(stats["tripDays"]),
                          "vehicleTraces": len(stats["traces"])}
                         for chain_id, stats in sorted(chain_stats.items())]}
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: analyze-kawasaki-position-history.py WORKBOOK STATIC")
    main(sys.argv[1], sys.argv[2])
